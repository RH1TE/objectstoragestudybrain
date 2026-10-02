from __future__ import annotations

import csv
import json
from pathlib import Path

import pymupdf
from docx import Document
from markdownify import markdownify
from openpyxl import load_workbook
from pptx import Presentation
from striprtf.striprtf import rtf_to_text

SUPPORTED = {
    ".pdf", ".docx", ".pptx", ".txt", ".md", ".rst", ".tex",
    ".html", ".htm", ".rtf", ".ipynb", ".csv", ".xlsx",
}


def clean_text(value: str) -> str:
    return (
        (value or "")
        .replace("\x00", "")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .strip()
    )


def convert_pdf(path: Path) -> tuple[str, str]:
    doc = pymupdf.open(path)
    output: list[str] = []
    extracted = 0
    blank_pages = 0

    for number, page in enumerate(doc, 1):
        text = clean_text(page.get_text("text", sort=True))
        extracted += len(text)
        output.append(f"## Page {number}\n")
        if text:
            output.append(text + "\n")
        else:
            blank_pages += 1
            output.append("[No extractable text on this page.]\n")

    if extracted == 0:
        status = "needs_ocr"
    elif blank_pages:
        status = "partial"
    else:
        status = "ok"

    if blank_pages:
        output.append(
            f"\n> {blank_pages} page(s) had no extractable text. "
            "The source may contain scans or images.\n"
        )
    return "\n".join(output), status


def convert_docx(path: Path) -> tuple[str, str]:
    doc = Document(path)
    output: list[str] = []

    for paragraph in doc.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading"):
            digits = "".join(ch for ch in style if ch.isdigit())
            level = max(2, min(6, int(digits or "2") + 1))
            output.append("#" * level + " " + text)
        else:
            output.append(text)

    for number, table in enumerate(doc.tables, 1):
        output.append(f"\n### Table {number}")
        for row in table.rows:
            cells = [clean_text(cell.text).replace("|", "\\|") for cell in row.cells]
            output.append("| " + " | ".join(cells) + " |")

    return "\n\n".join(output), "ok"


def convert_pptx(path: Path) -> tuple[str, str]:
    presentation = Presentation(path)
    output: list[str] = []

    for number, slide in enumerate(presentation.slides, 1):
        output.append(f"## Slide {number}")
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = clean_text(shape.text)
                if text:
                    parts.append(text)
        output.append("\n\n".join(parts) if parts else "[No extractable slide text.]")

    return "\n\n".join(output), "ok"

def convert_xlsx(path: Path) -> tuple[str, str]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    output: list[str] = []

    for sheet in workbook.worksheets:
        output.append(f"## Sheet: {sheet.title}")
        rows = []
        for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                rows.append(values)
            if number >= 5000:
                rows.append(["[truncated after 5000 rows]"])
                break

        width = max((len(row) for row in rows), default=0)
        for row in rows:
            row += [""] * (width - len(row))
            output.append(
                "| "
                + " | ".join(value.replace("|", "\\|").replace("\n", " ") for value in row)
                + " |"
            )

    return "\n".join(output), "ok"


def convert_csv(path: Path) -> tuple[str, str]:
    output: list[str] = []
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        for number, row in enumerate(csv.reader(handle), 1):
            output.append(
                "| "
                + " | ".join(value.replace("|", "\\|").replace("\n", " ") for value in row)
                + " |"
            )
            if number >= 5000:
                output.append("| [truncated after 5000 rows] |")
                break
    return "\n".join(output), "ok"

def convert_ipynb(path: Path) -> tuple[str, str]:
    notebook = json.loads(path.read_text(errors="replace"))
    output: list[str] = []

    for number, cell in enumerate(notebook.get("cells") or [], 1):
        source = "".join(cell.get("source") or []).strip()
        if not source:
            continue
        if cell.get("cell_type") == "markdown":
            output.append(f"## Notebook cell {number}\n\n{source}")
        else:
            output.append(
                f"## Notebook code cell {number}\n\n"
                f"```python\n{source}\n```"
            )

    return "\n\n".join(output), "ok"


def convert(path: Path) -> tuple[str, str]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return convert_pdf(path)
    if ext == ".docx":
        return convert_docx(path)
    if ext == ".pptx":
        return convert_pptx(path)
    if ext == ".xlsx":
        return convert_xlsx(path)
    if ext == ".csv":
        return convert_csv(path)
    if ext == ".ipynb":
        return convert_ipynb(path)
    if ext in {".html", ".htm"}:
        return clean_text(markdownify(path.read_text(errors="replace"), heading_style="ATX")), "ok"
    if ext == ".rtf":
        return clean_text(rtf_to_text(path.read_text(errors="replace"))), "ok"

    text = clean_text(path.read_text(errors="replace"))
    if ext == ".tex":
        return f"```latex\n{text}\n```", "ok"
    return text, "ok"
