from pathlib import Path

import pymupdf
from docx import Document

from study_brain.converters import convert


def test_pdf_keeps_page_markers(tmp_path: Path):
    path = tmp_path / "notes.pdf"
    document = pymupdf.open()
    first = document.new_page()
    first.insert_text((72, 72), "Wave optics and diffraction")
    second = document.new_page()
    second.insert_text((72, 72), "Interference and phase")
    document.save(path)

    text, status = convert(path)

    assert status == "ok"
    assert "## Page 1" in text
    assert "## Page 2" in text
    assert "Wave optics and diffraction" in text


def test_docx_extracts_headings_and_text(tmp_path: Path):
    path = tmp_path / "workshop.docx"
    document = Document()
    document.add_heading("Workshop", level=1)
    document.add_paragraph("A short worked example.")
    document.save(path)

    text, status = convert(path)

    assert status == "ok"
    assert "Workshop" in text
    assert "A short worked example." in text
